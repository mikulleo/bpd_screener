import io

import openpyxl
from fastapi.testclient import TestClient

from screener.config import Thresholds
from screener.server import app
from screener.tickers import guess_phase, note_date, parse_file, parse_text

client = TestClient(app)


def test_parse_text_handles_watchlist_quirks():
    text = "Ticker\nADBE\nbrk/b\nNWSA / NWS\nAAPL, msft, BRK.A, GOOGL/GOOG\n$SPY\nADBE\n"
    assert parse_text(text) == ["ADBE", "BRK-B", "NWSA", "NWS", "AAPL", "MSFT", "BRK-A", "GOOGL", "GOOG", "SPY"]


def test_parse_xlsx_reads_ticker_column_and_metadata():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S&P500"
    ws.append(["Ticker", "Company", "Industry", "Signal", None, "Leos Review", None])
    ws.append(["ADBE", "Adobe\xa0", "Software", "BUY", None, "OK", "watch earnings"])
    ws.append(["NWSA / NWS", "News Corp", "Media", "PERFORMER", None, "OK, not prio", None])
    ws2 = wb.create_sheet("2b AB")
    ws2.append(["Fáze", "Ticker", "Název", "Poznámka"])
    ws2.append(["2b AB", "AEIS", "Advanced Energy", "16/8 LM: Presunuto z 2b"])
    ws3 = wb.create_sheet("Notes")
    ws3.append(["Reminder of the concepts:"])
    buf = io.BytesIO()
    wb.save(buf)
    sheets = parse_file("wl.xlsx", buf.getvalue())
    assert [s["name"] for s in sheets] == ["S&P500", "2b AB"]
    sp = sheets[0]
    assert sp["tickers"] == ["ADBE", "NWSA", "NWS"]
    assert sp["meta"]["ADBE"] == {"name": "Adobe", "industry": "Software", "tags": "BUY", "note": "OK — watch earnings"}
    assert sp["phase"] is None
    assert sheets[1]["phase"] == "2b-AB"
    assert sheets[1]["meta"]["AEIS"]["note"].startswith("16/8 LM")


def test_guess_phase_from_sheet_names():
    assert guess_phase("Minervini 2b+ RS+") == "2b+RS+"
    assert guess_phase("Minervini 2b+ RS-") == "2b+RS-"
    assert guess_phase("Weinstein 2a+") == "2a+"
    assert guess_phase("2b") == "2b"
    assert guess_phase("3") == "3"
    assert guess_phase("S&P500") is None


def test_note_date_picks_latest_and_rolls_back_future_dates():
    from datetime import datetime, timezone
    today = datetime(2026, 10, 23, tzinfo=timezone.utc)
    assert note_date("20/7 LM: korekce | 16/8 LM: obrat", today).date().isoformat() == "2026-08-16"
    assert note_date("14/11 LM: x", today).year == 2025
    assert note_date("no date here", today) is None


def test_threshold_overrides_cast_and_ignore_unknown():
    t = Thresholds.from_overrides({"rs_pass_pct": "3", "base_pass_weeks": 10.0, "bogus": 1})
    assert t.rs_pass_pct == 3.0 and t.base_pass_weeks == 10


def test_password_protects_api(monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "secret")
    assert client.get("/api/config").json()["auth_required"] is True
    assert client.get("/api/state").status_code == 401
    assert client.get("/api/auth-check", headers={"X-App-Password": "secret"}).status_code == 200
    assert client.post("/api/evaluate", json={"symbols": ["AAPL"]}).status_code == 401


def test_evaluate_rejects_oversized_batches():
    r = client.post("/api/evaluate", json={"symbols": [f"T{i}" for i in range(30)]})
    assert r.status_code == 400


def test_lists_notes_and_manual_phase_roundtrip():
    lid = client.post("/api/lists", json={"name": "Core"}).json()["id"]
    r = client.post(f"/api/lists/{lid}/add", json={"symbols": ["aapl", "brk/b"]}).json()
    assert r["added"] == ["AAPL", "BRK-B"]
    assert next(li for li in r["lists"] if li["id"] == lid)["count"] == 2

    row = client.patch("/api/tickers/AAPL", json={"note": "leader", "manual_phase": "2b-AB"}).json()
    assert row["note"] == "leader" and row["manual_phase"] == "2b-AB" and row["manual_phase_source"] == "user"
    assert client.patch("/api/tickers/AAPL", json={"manual_phase": "5"}).status_code == 400

    state = client.get("/api/state").json()
    aapl = next(t for t in state["tickers"] if t["symbol"] == "AAPL")
    assert aapl["lists"] == [lid]
    assert any(lab["symbol"] == "AAPL" and lab["phase"] == "2b-AB" for lab in client.get("/api/labels").json()["labels"])

    client.post(f"/api/lists/{lid}/remove", json={"symbols": ["AAPL"]})
    assert next(li for li in client.get("/api/state").json()["lists"] if li["id"] == lid)["count"] == 1
    client.delete(f"/api/lists/{lid}")
    assert all(li["id"] != lid for li in client.get("/api/state").json()["lists"])


def test_import_commit_creates_lists_and_dated_manual_phases():
    body = {"sheets": [
        {"list_name": "Midcaps", "tickers": ["BILL"], "meta": {"BILL": {"name": "Bill Holdings", "industry": "Software"}}},
        {"phase": "2b-AB", "tickers": ["BILL"], "meta": {"BILL": {"note": "20/7 LM: z 2b"}}},
    ]}
    assert client.post("/api/import/commit", json=body).json()["tickers"] == 1
    bill = next(t for t in client.get("/api/state").json()["tickers"] if t["symbol"] == "BILL")
    assert bill["name"] == "Bill Holdings" and bill["manual_phase"] == "2b-AB"
    assert bill["manual_phase_source"] == "import" and "-07-20" in bill["manual_phase_at"]


def test_thresholds_are_shared_settings():
    client.put("/api/settings/thresholds", json={"rs_pass_pct": 4, "base_pass_weeks": 14, "nope": 1})
    assert client.get("/api/state").json()["thresholds"] == {"rs_pass_pct": 4}


def test_static_index_served():
    r = client.get("/")
    assert r.status_code == 200 and "BPD Watchlist" in r.text
